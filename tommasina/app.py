# Tommasina — Quote-part de terrain et chiffre d'affaires par lot.
# Lancer avec : python app.py   (puis http://127.0.0.1:8001)
# Les données sont dans data.db (SQLite) : elles sont conservées d'un lancement à l'autre.

import io
import json
import shutil
import sqlite3
import tempfile
from datetime import date, datetime
from pathlib import Path
from typing import Optional

import openpyxl
import uvicorn
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel

BASE = Path(__file__).parent
DB_PATH = BASE / "data.db"
SEED = BASE / "seed.json"

CHAMPS = ["date_contrat", "client", "tf", "propriete", "indice", "niveau", "consistance",
          "surface_titre", "surface_totale", "prix_m2", "prix_ttc", "encaisse", "notes"]
PARAMS_DEFAUT = {"surface_terrain": 692, "prix_terrain_total": 12812552.5, "prix_m2": 18515.25, "tva": 0.2}


def connexion():
    c = sqlite3.connect(DB_PATH)
    c.row_factory = sqlite3.Row
    return c


def init_db():
    with connexion() as c:
        c.executescript("""
        CREATE TABLE IF NOT EXISTS lots (
            id INTEGER PRIMARY KEY AUTOINCREMENT, ordre INTEGER NOT NULL DEFAULT 0,
            date_contrat TEXT, client TEXT, tf TEXT, propriete TEXT, indice TEXT, niveau TEXT,
            consistance TEXT, surface_titre REAL, surface_totale REAL, prix_m2 REAL,
            prix_ttc REAL, encaisse REAL, notes TEXT);
        CREATE TABLE IF NOT EXISTS parametres (cle TEXT PRIMARY KEY, valeur REAL NOT NULL);
        """)
        if c.execute("SELECT COUNT(*) FROM parametres").fetchone()[0] == 0:
            donnees = json.loads(SEED.read_text(encoding="utf-8")) if SEED.exists() else {"parametres": PARAMS_DEFAUT, "lots": []}
            c.executemany("INSERT INTO parametres VALUES (?,?)", donnees["parametres"].items())
            for i, lot in enumerate(donnees["lots"]):
                inserer(c, lot, i)


def inserer(c, lot: dict, ordre: int):
    cur = c.execute(f"INSERT INTO lots (ordre,{','.join(CHAMPS)}) VALUES (?,{','.join('?' * len(CHAMPS))})",
                    [ordre] + [lot.get(k) for k in CHAMPS])
    return cur.lastrowid


def parametres(c) -> dict:
    p = dict(PARAMS_DEFAUT)
    p.update({r["cle"]: r["valeur"] for r in c.execute("SELECT * FROM parametres")})
    return p


def calculer(lot: dict, p: dict) -> dict:
    """Mêmes formules que le tableau Excel."""
    out = dict(lot)
    S = lot.get("surface_totale")
    prix_m2 = lot.get("prix_m2") or p["prix_m2"]
    out["prix_m2_effectif"] = prix_m2
    out["quote_part_m2"] = S * p["surface_terrain"] / 10000 if S else None          # J
    out["part_terrain"] = out["quote_part_m2"] * prix_m2 if S else None              # M
    ttc = lot.get("prix_ttc")
    if ttc and S:
        out["prix_hors_terrain"] = ttc - out["part_terrain"]                         # N
        out["prix_ht"] = out["prix_hors_terrain"] / (1 + p["tva"])                   # O
        out["tva"] = out["prix_ht"] * p["tva"]                                       # P
    else:
        out["prix_hors_terrain"] = out["prix_ht"] = out["tva"] = None
    out["vendu"] = bool(ttc)
    out["taux_encaisse"] = lot["encaisse"] / ttc if lot.get("encaisse") and ttc else None
    return out


def tous_les_lots(c):
    p = parametres(c)
    return [calculer(dict(r), p) for r in c.execute("SELECT * FROM lots ORDER BY ordre, id")], p


class LotIn(BaseModel):
    date_contrat: Optional[str] = None
    client: Optional[str] = None
    tf: Optional[str] = None
    propriete: Optional[str] = None
    indice: Optional[str] = None
    niveau: Optional[str] = None
    consistance: Optional[str] = None
    surface_titre: Optional[float] = None
    surface_totale: Optional[float] = None
    prix_m2: Optional[float] = None
    prix_ttc: Optional[float] = None
    encaisse: Optional[float] = None
    notes: Optional[str] = None


def nettoyer(d: dict) -> dict:
    for k, v in d.items():
        if isinstance(v, str):
            d[k] = v.strip() or None
    if d.get("date_contrat"):
        try:
            date.fromisoformat(d["date_contrat"])
        except ValueError:
            raise HTTPException(422, "Date invalide (format AAAA-MM-JJ)")
    return d


app = FastAPI(title="Tommasina")


@app.on_event("startup")
def _startup():
    init_db()


@app.get("/")
def accueil():
    return FileResponse(BASE / "static" / "index.html")


@app.get("/api/lots")
def api_lots():
    with connexion() as c:
        lots, p = tous_les_lots(c)
    return {"lots": lots, "parametres": p}


@app.post("/api/lots")
def api_creer(lot: LotIn):
    with connexion() as c:
        ordre = c.execute("SELECT COALESCE(MAX(ordre),0)+1 FROM lots").fetchone()[0]
        i = inserer(c, nettoyer(lot.model_dump()), ordre)
        return calculer(dict(c.execute("SELECT * FROM lots WHERE id=?", (i,)).fetchone()), parametres(c))


@app.put("/api/lots/{lot_id}")
def api_modifier(lot_id: int, lot: LotIn):
    with connexion() as c:
        if not c.execute("SELECT 1 FROM lots WHERE id=?", (lot_id,)).fetchone():
            raise HTTPException(404, "Lot introuvable")
        d = nettoyer(lot.model_dump())
        c.execute(f"UPDATE lots SET {','.join(k + '=?' for k in CHAMPS)} WHERE id=?", [d[k] for k in CHAMPS] + [lot_id])
        return calculer(dict(c.execute("SELECT * FROM lots WHERE id=?", (lot_id,)).fetchone()), parametres(c))


@app.delete("/api/lots/{lot_id}")
def api_supprimer(lot_id: int):
    with connexion() as c:
        c.execute("DELETE FROM lots WHERE id=?", (lot_id,))
    return {"ok": True}


class ParamsIn(BaseModel):
    surface_terrain: float
    prix_terrain_total: float
    prix_m2: float
    tva: float


@app.put("/api/parametres")
def api_parametres(p: ParamsIn):
    if p.surface_terrain <= 0 or not 0 <= p.tva < 1:
        raise HTTPException(422, "Paramètres invalides")
    with connexion() as c:
        c.executemany("INSERT OR REPLACE INTO parametres VALUES (?,?)", p.model_dump().items())
    return api_lots()


# ---- Export / import / sauvegarde -------------------------------------------------------------

COLONNES_XLSX = [
    ("Date Contrat", "date_contrat"), ("Client", "client"), ("TF", "tf"), ("Propriété", "propriete"),
    ("Indice de la partie privative", "indice"), ("Niveau", "niveau"), ("Consistance", "consistance"),
    ("A l'intérieur du titre (m²)", "surface_titre"), ("Total avec surplomb (m²)", "surface_totale"),
    ("Quote-part terrain m²", "quote_part_m2"), ("Prix m²", "prix_m2_effectif"),
    ("Total Dhs part terrain", "part_terrain"), ("Prix de vente hors part terrain", "prix_hors_terrain"),
    ("Prix de vente HT", "prix_ht"), ("TVA", "tva"), ("TTC", "prix_ttc"), ("Encaissé", "encaisse"),
    ("Notes", "notes"),
]


@app.get("/api/export.xlsx")
def export_xlsx():
    with connexion() as c:
        lots, _ = tous_les_lots(c)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tommasina"
    ws.append([t for t, _ in COLONNES_XLSX])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3A5F")
    for lot in lots:
        row = []
        for _, k in COLONNES_XLSX:
            v = lot.get(k)
            row.append(date.fromisoformat(v) if k == "date_contrat" and v else v)
        ws.append(row)
    n = len(lots) + 1
    ws.append([])
    ws.append(["TOTAL"] + [None] * 6 + [f"=SUM({col}2:{col}{n})" if col else None
                                          for col in "HIJ"] + [None] + [f"=SUM({col}2:{col}{n})" for col in "LMNOPQ"])
    for r in ws.iter_rows(min_row=2, max_col=1):
        r[0].number_format = "dd/mm/yyyy"
    for r in ws.iter_rows(min_row=2, min_col=10, max_col=17):
        for cell in r:
            cell.number_format = "#,##0.00"
    for i, w in enumerate([12, 30, 14, 14, 12, 14, 14, 10, 10, 12, 12, 16, 18, 16, 14, 14, 14, 30], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "E2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="Tommasina_lots.xlsx"'})


@app.get("/api/sauvegarde")
def sauvegarde():
    """Copie cohérente de la base (à conserver comme sauvegarde)."""
    tmp = Path(tempfile.mkdtemp()) / "sauvegarde.db"
    with connexion() as c:
        dest = sqlite3.connect(tmp)
        c.backup(dest)
        dest.close()
    nom = f"tommasina_sauvegarde_{datetime.now():%Y%m%d_%H%M}.db"
    return FileResponse(tmp, media_type="application/octet-stream", filename=nom)


@app.post("/api/restaurer")
async def restaurer(fichier: UploadFile = File(...)):
    contenu = await fichier.read()
    tmp = Path(tempfile.mkdtemp()) / "restore.db"
    tmp.write_bytes(contenu)
    try:
        t = sqlite3.connect(tmp)
        t.execute("SELECT COUNT(*) FROM lots")
        t.execute("SELECT COUNT(*) FROM parametres")
        t.close()
    except sqlite3.Error:
        raise HTTPException(422, "Ce fichier n'est pas une sauvegarde Tommasina valide")
    if DB_PATH.exists():
        shutil.copy(DB_PATH, DB_PATH.with_suffix(".avant-restauration.db"))
    shutil.copy(tmp, DB_PATH)
    return {"ok": True}


if __name__ == "__main__":
    init_db()
    uvicorn.run(app, host="127.0.0.1", port=8001)
